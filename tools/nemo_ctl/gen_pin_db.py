#!/usr/bin/env python3
"""Generate data/pin_db.yaml from the lab's NEMO chip workbooks.

Reads docs/reference/nemo_chip/Blocks.xlsx ("Blocks Interface - Digital" sheet) and
docs/reference/nemo_chip/IO_pinout.xlsx ("IO-Pins-Dima" sheet), and applies the
reconciliation worked out in docs/controller_spec.md (sections 2 and 3) to produce the
union pin list: 22 inputs + 4 outputs, with both sheets' pad numbers kept side by side
rather than silently picking one where they disagree.

The per-pin safe limits / idle level / polarity / notes are NOT in either workbook (they
don't exist there); those come from docs/controller_spec.md section 3 and are encoded
below in PIN_METADATA, each one traceable to an open question (Qn) if unconfirmed.

Re-run this after the workbooks change:
    python3 tools/nemo_ctl/gen_pin_db.py
It will raise if a name this script expects can no longer be found, so drift in the
workbooks is caught rather than silently producing a stale pin_db.yaml.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import yaml

REPO_ROOT = Path(__file__).resolve().parents[2]
BLOCKS_XLSX = REPO_ROOT / "docs/reference/nemo_chip/Blocks.xlsx"
IO_PINOUT_XLSX = REPO_ROOT / "docs/reference/nemo_chip/IO_pinout.xlsx"
OUT_PATH = Path(__file__).resolve().parent / "data" / "pin_db.yaml"

DIGITAL_SHEET_NAME = "Blocks Interface - Digital"
IO_PINS_SHEET_NAME = "IO-Pins-Dima"


def _load_digital_sheet_pads(path: Path) -> dict[str, int]:
    """name -> pad number, from the 1-68 physical pad-ring sheet. First match wins;
    duplicate names (vssx, Unused) are not looked up by canonical pins so collisions
    there don't matter."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[DIGITAL_SHEET_NAME]
    pads: dict[str, int] = {}
    for row in ws.iter_rows(values_only=True):
        pad, name = row[0], row[1]
        if isinstance(pad, int) and isinstance(name, str) and name not in pads:
            pads[name] = pad
    return pads


def _load_io_pins_dima_pads(path: Path) -> dict[str, int]:
    """name -> pin number, from IO-Pins-Dima's linear 1-67 list (the table at the top
    of the sheet, not the pictorial pad-ring diagram further down — see
    docs/controller_spec.md §2.1, that diagram's numbering was not extracted here)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[IO_PINS_SHEET_NAME]
    pads: dict[str, int] = {}
    for row in ws.iter_rows(values_only=True):
        pin, name = row[0], row[1]
        if isinstance(pin, int) and isinstance(name, str) and name not in pads:
            pads[name] = pin
    return pads


# Canonical pin -> name as it appears on each sheet (None = absent from that sheet).
# This *is* the reconciliation from docs/controller_spec.md §2/§3, made explicit so a
# re-run can verify it against the live workbooks rather than re-deriving it by hand.
# fmt: off
CANONICAL_PINS = [
    # name,          block,                direction, digital_sheet_name, io_pins_dima_name
    ("clk",          "Clock Generator",     "in",  "clk",       "clk"),
    ("reset",        "Clock Generator",     "in",  "reset",     "reset"),
    ("set",          "Clock Generator",     "in",  "set",       "set"),
    ("rst_dsa",      "DSA",                 "in",  "rst_dsa",   "rst_dsa"),
    ("pulse",        "Array with Muxes",    "in",  "pulse",     "pulse"),
    ("s0_p",         "Array with Muxes",    "in",  "s0_p",      "s0"),
    ("s1_p",         "Array with Muxes",    "in",  "s1_p",      "s1"),
    ("s2_p",         "Array with Muxes",    "in",  "s2_p",      "s2"),
    ("b0_dft",       "DFT - All",           "in",  "b0_dft",    "b0_dft"),
    ("b1_dft",       "DFT - All",           "in",  "b1_dft",    "b1_dft"),
    ("b2_dft",       "DFT - All",           "in",  "b2_dft",    "b2_dft"),
    ("s0_dft",       "dft per ADC",         "in",  "s0_dft",    "s0_dft"),
    ("s1_dft",       "dft per ADC",         "in",  "s1_dft",    "s1_dft"),
    ("s2_dft",       "dft per ADC",         "in",  "s2_dft",    None),
    ("s3_dft",       "dft per ADC",         "in",  "s3_dft",    None),
    ("sipo_in",      "SIPO",                "in",  "sipo_in",   "sipo_in"),
    ("clk_sipo",     "SIPO",                "in",  "clk_sipo",  "clk_sipo"),
    ("cs",           "SIPO",                "in",  "cs",        "cs"),
    ("rst_sipo",     "SIPO",                "in",  None,        "rst_sipo"),
    ("clk_piso",     "PISO",                "in",  "clk_piso",  "clk_piso"),
    ("shift",        "PISO",                "in",  "shift",     "shift"),
    ("rst_piso",     "PISO",                "in",  "rst_piso",  "rst_piso"),
    ("sipo_out",     "SIPO",                "out", "sipo_out",  "sipo_out"),
    ("piso_out",     "PISO",                "out", "piso_out",  "piso_out"),
    ("dft0_buff",    "DFT - All",           "out", "dft0_buff", "dft0"),
    ("dft1_buff",    "DFT - All",           "out", "dft1_buff", "dft1"),
]
# fmt: on

N_INPUTS = sum(1 for p in CANONICAL_PINS if p[2] == "in")
N_OUTPUTS = sum(1 for p in CANONICAL_PINS if p[2] == "out")

# Known, expected absences (docs/controller_spec.md §2.1, §2.2) — regeneration should
# NOT raise over these specific gaps, only over new/unexpected ones.
EXPECTED_ABSENT_FROM_DIGITAL_SHEET = {"rst_sipo"}
EXPECTED_ABSENT_FROM_IO_PINS_DIMA = {"s2_dft", "s3_dft"}

# Per-pin metadata that doesn't exist in either workbook — from docs/controller_spec.md
# §3. allowed_types uses the 7 generic types from generic_controller_prompt.md §1.
# limits/idle_level/polarity are left null (TBD) unless the RTL or the explainer actually
# states them; "notes" cites the open question number (Qn) from controller_spec.md §5
# where the value is still unconfirmed.
PIN_METADATA: dict[str, dict] = {
    "clk": dict(allowed_types=["CLOCK"], limits={"max_freq_hz": None}, idle_level="LOW",
                polarity=None,
                notes="Drives on-chip phi1/phi2/phiR generation; max rate TBD (Q7)."),
    "reset": dict(allowed_types=["LOW", "HIGH", "PULSE"], limits={}, idle_level=None,
                  polarity=None, notes="Meaning/polarity undocumented (Q6)."),
    "set": dict(allowed_types=["LOW", "HIGH", "PULSE"], limits={}, idle_level=None,
                polarity=None, notes="Meaning/polarity undocumented (Q6)."),
    "rst_dsa": dict(allowed_types=["LOW", "HIGH", "PULSE"], limits={}, idle_level=None,
                    polarity=None, notes="Clears DSA accumulators; polarity TBD (Q6)."),
    "pulse": dict(allowed_types=["PULSE"], limits={"min_width_ticks": None}, idle_level=None,
                  polarity=None,
                  notes="Probably program/erase pulse timing; width/meaning TBD (Q6, Q7)."),
    "s0_p": dict(allowed_types=["LOW", "HIGH", "BUS"], limits={}, idle_level=None,
                 polarity=None, notes="3-bit mode/op select with s1_p,s2_p; meaning TBD (Q6)."),
    "s1_p": dict(allowed_types=["LOW", "HIGH", "BUS"], limits={}, idle_level=None,
                 polarity=None, notes="See s0_p."),
    "s2_p": dict(allowed_types=["LOW", "HIGH", "BUS"], limits={}, idle_level=None,
                 polarity=None, notes="See s0_p."),
    "b0_dft": dict(allowed_types=["LOW", "HIGH", "BUS"], limits={}, idle_level=None,
                   polarity=None, notes="DFT-All pair select; lookup table UNCONFIRMED (spec §4)."),
    "b1_dft": dict(allowed_types=["LOW", "HIGH", "BUS"], limits={}, idle_level=None,
                   polarity=None, notes="See b0_dft."),
    "b2_dft": dict(allowed_types=["LOW", "HIGH", "BUS"], limits={}, idle_level=None,
                   polarity=None, notes="See b0_dft."),
    "s0_dft": dict(allowed_types=["LOW", "HIGH", "BUS"], limits={}, idle_level=None,
                   polarity=None,
                   notes="dft-per-ADC d0..d7 select; bit count disagrees between sheets (Q10)."),
    "s1_dft": dict(allowed_types=["LOW", "HIGH", "BUS"], limits={}, idle_level=None,
                   polarity=None, notes="See s0_dft."),
    "s2_dft": dict(allowed_types=["LOW", "HIGH", "BUS"], limits={}, idle_level=None,
                   polarity=None,
                   notes="Absent from IO-Pins-Dima's linear list (Q10) — exists per the digital sheet."),
    "s3_dft": dict(allowed_types=["LOW", "HIGH", "BUS"], limits={}, idle_level=None,
                   polarity=None,
                   notes="Absent from IO-Pins-Dima's linear list (Q10) — exists per the digital sheet."),
    "sipo_in": dict(allowed_types=["SERIAL"], limits={"frame_bits": 136}, idle_level="LOW",
                    polarity=None,
                    notes="MSB first, bit135 first; must change on clk_sipo falling edge."),
    "clk_sipo": dict(allowed_types=["CLOCK"], limits={"max_freq_hz": None, "cycles": 136},
                      idle_level="LOW", polarity=None,
                      notes="136 edges then stop low; max rate TBD (Q7)."),
    "cs": dict(allowed_types=["PULSE"], limits={"min_width_ticks": 1}, idle_level="LOW",
               polarity="active_high",
               notes="Edge-triggered, not a level; must come from a registered FF."),
    "rst_sipo": dict(allowed_types=["LOW", "HIGH", "PULSE"], limits={}, idle_level="HIGH",
                      polarity="active_low",
                      notes="Active-low confirmed from RTL (rst_n); pad-level polarity TBD (Q6). "
                            "Pad number itself unconfirmed: 37 (IO-Pins-Dima list) vs 39 "
                            "(brief's note on the floorplan diagram) — see spec §2.1."),
    "clk_piso": dict(allowed_types=["CLOCK"], limits={"max_freq_hz": None, "cycles": 168},
                      idle_level="LOW", polarity=None,
                      notes="1 load edge + 167 shift edges; max rate TBD (Q7)."),
    "shift": dict(allowed_types=["LOW", "HIGH"], limits={}, idle_level="LOW", polarity=None,
                   notes="0=load/sample, 1=shift (RTL `load`); sequencing is a Phase-7 concern."),
    "rst_piso": dict(allowed_types=["LOW", "HIGH", "PULSE"], limits={}, idle_level="HIGH",
                      polarity="active_low",
                      notes="Active-low confirmed from RTL (rst_n); pad-level polarity TBD (Q6)."),
    "sipo_out": dict(allowed_types=["CAPTURE"], limits={}, idle_level=None, polarity=None,
                      notes="shift_reg[0], one-clock-late echo of sipo_in; liveness check only."),
    "piso_out": dict(allowed_types=["CAPTURE"], limits={"frame_bits": 168}, idle_level=None,
                      polarity=None,
                      notes="MSB first, split 4x[14+7+14+7]; bit 1 valid before any shift clocks, "
                            "sample thereafter on clk_piso falling edge."),
    "dft0_buff": dict(allowed_types=["CAPTURE"], limits={}, idle_level=None, polarity=None,
                       notes="Buffered DFT observation output 0; selection via DFT helper (UNCONFIRMED)."),
    "dft1_buff": dict(allowed_types=["CAPTURE"], limits={}, idle_level=None, polarity=None,
                       notes="Buffered DFT observation output 1; selection via DFT helper (UNCONFIRMED)."),
}


def build_pin_db() -> list[dict]:
    digital_pads = _load_digital_sheet_pads(BLOCKS_XLSX)
    io_pins_pads = _load_io_pins_dima_pads(IO_PINOUT_XLSX)

    entries = []
    for name, block, direction, digi_name, iopd_name in CANONICAL_PINS:
        pad_digital = digital_pads.get(digi_name) if digi_name else None
        pad_iopd = io_pins_pads.get(iopd_name) if iopd_name else None

        if digi_name and pad_digital is None and name not in EXPECTED_ABSENT_FROM_DIGITAL_SHEET:
            raise RuntimeError(
                f"{name!r}: expected to find {digi_name!r} in {DIGITAL_SHEET_NAME!r} "
                f"but it's gone — workbook changed, re-check the reconciliation."
            )
        if iopd_name and pad_iopd is None and name not in EXPECTED_ABSENT_FROM_IO_PINS_DIMA:
            raise RuntimeError(
                f"{name!r}: expected to find {iopd_name!r} in {IO_PINS_SHEET_NAME!r} "
                f"but it's gone — workbook changed, re-check the reconciliation."
            )

        meta = PIN_METADATA[name]
        entries.append({
            "name": name,
            "block": block,
            "direction": direction,
            "pad_digital_sheet": pad_digital,
            "pad_io_pins_dima": pad_iopd,
            "fmc_la": None,  # TBD for every pin — see controller_spec.md §3
            "allowed_types": meta["allowed_types"],
            "limits": meta["limits"],
            "idle_level": meta["idle_level"],
            "polarity": meta["polarity"],
            "notes": meta["notes"],
        })

    if sum(1 for e in entries if e["direction"] == "in") != N_INPUTS:
        raise RuntimeError("input count drifted from CANONICAL_PINS")
    if sum(1 for e in entries if e["direction"] == "out") != N_OUTPUTS:
        raise RuntimeError("output count drifted from CANONICAL_PINS")
    return entries


def main() -> None:
    entries = build_pin_db()
    doc = {
        "_generated_by": "tools/nemo_ctl/gen_pin_db.py",
        "_source": [
            "docs/reference/nemo_chip/Blocks.xlsx!Blocks Interface - Digital",
            "docs/reference/nemo_chip/IO_pinout.xlsx!IO-Pins-Dima",
            "docs/controller_spec.md (reconciliation + per-pin metadata)",
        ],
        "_count": {"inputs": N_INPUTS, "outputs": N_OUTPUTS, "total": len(entries)},
        "pins": entries,
    }
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with OUT_PATH.open("w") as f:
        yaml.safe_dump(doc, f, sort_keys=False, width=100)
    print(f"wrote {OUT_PATH} ({len(entries)} pins: {N_INPUTS} in + {N_OUTPUTS} out)")


if __name__ == "__main__":
    try:
        main()
    except RuntimeError as exc:
        print(f"error: {exc}", file=sys.stderr)
        sys.exit(1)
