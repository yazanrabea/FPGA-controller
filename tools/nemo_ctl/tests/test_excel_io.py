from pathlib import Path

from nemo_ctl.config import config_to_dict, load_config
from nemo_ctl.excel_io import export_config, export_template, import_config
from nemo_ctl.pins import load_pin_db

EXAMPLE_CONFIG = Path(__file__).resolve().parents[3] / "config" / "nemo_config.yaml"


def test_export_import_round_trip(tmp_path):
    cfg = load_config(EXAMPLE_CONFIG)
    xlsx = tmp_path / "roundtrip.xlsx"
    export_config(cfg, xlsx)
    cfg2 = import_config(xlsx)
    assert config_to_dict(cfg)["modes"] == config_to_dict(cfg2)["modes"]


def test_exported_workbook_has_one_sheet_per_mode(tmp_path):
    cfg = load_config(EXAMPLE_CONFIG)
    xlsx = tmp_path / "modes.xlsx"
    export_config(cfg, xlsx)
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    assert set(wb.sheetnames) == set(cfg.mode_names())


def test_template_has_one_row_per_input_pin_per_mode(tmp_path):
    db = load_pin_db()
    xlsx = tmp_path / "template.xlsx"
    export_template(db, xlsx, modes=("INFERENCE",))
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ws = wb["INFERENCE"]
    data_rows = list(ws.iter_rows(values_only=True))[1:]
    n_inputs = sum(1 for p in db.values() if p.direction == "in")
    assert len(data_rows) == n_inputs
